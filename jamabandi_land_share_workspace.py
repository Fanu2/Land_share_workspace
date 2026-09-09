import sys
import csv
import webbrowser
from fractions import Fraction

from PySide6.QtCore import Qt
from PySide6.QtGui import QAction
from PySide6.QtWidgets import (
    QApplication, QAbstractItemView, QComboBox, QFileDialog, QFrame,
    QGridLayout, QGroupBox, QHBoxLayout, QLabel, QLineEdit, QMainWindow,
    QMessageBox, QPushButton, QSplitter, QTableWidget, QTableWidgetItem,
    QVBoxLayout, QWidget, QHeaderView
)

DATA = [(1, 11, 'जसवीर सिंह', 'गुरबचन सिंह', 'विशन सिंह', '30/413', 20, 13, '1K-10M-0S', 'Bhambu', 'https://taruana-archive.vercel.app/2023-24/1.htm'), (596, 849, 'जसवीर सिंह', 'गुरबचन सिंह', 'विशन सिंह', '1/1', 16, 13, '16K-13M-0S', 'behlan', 'https://taruana-archive.vercel.app/2023-24/596.htm'), (379, 579, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '9662/1003227', 380, 12, '3K-13M-3S', 'Geba', 'https://taruana-archive.vercel.app/2023-24/379.htm'), (379, 580, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '1/15', 19, 11, '1K-6M-1S', 'school', 'https://taruana-archive.vercel.app/2023-24/379.htm'), (388, 590, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '1/3', 61, 5, '20K-8M-3S', 'Bhatha', 'https://taruana-archive.vercel.app/2023-24/388.htm'), (393, 595, 'जसवीर सिंह', 'गुरबचन सिंह', 'विशन सिंह', '1/1', 85, 3, '85K-3M-0S', 'Boote wala', 'https://taruana-archive.vercel.app/2023-24/393.htm'), (464, 689, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '1/3', 12, 16, '4K-5M-3S', 'Bughe ali', 'https://taruana-archive.vercel.app/2023-24/464.htm'), (570, 810, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '2/769', 76, 18, '0K-4M-0S', 'sikh ali', 'https://taruana-archive.vercel.app/2023-24/570.htm'), (574, 814, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '100/1251', 13, 18, '1K-2M-2S', 'Kothe kole', 'https://taruana-archive.vercel.app/2023-24/574.htm'), (575, 815, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '97/1440', 32, 0, '2K-3M-1S', 'Dhani', 'https://taruana-archive.vercel.app/2023-24/575.htm'), (576, 816, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '2/69', 1, 3, '0K-0M-6S', 'Rasta behlan', 'https://taruana-archive.vercel.app/2023-24/576.htm'), (582, 823, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '6/96', 30, 14, '1K-18M-4S', 'Ram wali', 'https://taruana-archive.vercel.app/2023-24/582.htm'), (582, 823, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '2/96', 30, 14, '0K-12M-8S', 'Ram wali', 'https://taruana-archive.vercel.app/2023-24/582.htm'), (583, 824, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '353/19170', 100, 18, '1K-17M-2S', 'Sukhcain road', 'https://taruana-archive.vercel.app/2023-24/583.htm'), (583, 825, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '353/19170', None, 0, '0K-0M-0S', '', 'https://taruana-archive.vercel.app/2023-24/583.htm'), (583, 826, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '353/19170', None, 0, '0K-0M-0S', '', 'https://taruana-archive.vercel.app/2023-24/583.htm'), (583, 827, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '353/19170', None, 0, '0K-0M-0S', '', 'https://taruana-archive.vercel.app/2023-24/583.htm'), (583, 828, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '353/19170', None, 0, '0K-0M-0S', '', 'https://taruana-archive.vercel.app/2023-24/583.htm'), (584, 829, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '353/19170', 5, 12, '0K-2M-1S', 'Sukhcain road', 'https://taruana-archive.vercel.app/2023-24/584.htm'), (586, 831, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '1/3', 0, 16, '0K-5M-3S', 'khatti', 'https://taruana-archive.vercel.app/2023-24/586.htm'), (588, 833, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '7/198', 99, 0, '3K-10M-0S', '', 'https://taruana-archive.vercel.app/2023-24/588.htm'), (588, 834, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशન सिंह', '7/198', None, 0, '0K-0M-0S', '', 'https://taruana-archive.vercel.app/2023-24/588.htm'), (588, 835, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशન सिंह', '7/198', None, 0, '0K-0M-0S', '', 'https://taruana-archive.vercel.app/2023-24/588.htm'), (588, 836, 'जसवीર सिंह', 'गુરਬચਨ सिंह', 'बिशन सिंह', '7/198', None, 0, '0K-0M-0S', '', 'https://taruana-archive.vercel.app/2023-24/588.htm'), (652, 917, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '211/5811', 96, 17, '3K-10M-3S', 'Mana wali', 'https://taruana-archive.vercel.app/2023-24/652.htm'), (593, 845, 'जसवीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '69/367', 18, 7, '3K-9M-0S', 'behlan', 'https://taruana-archive.vercel.app/2023-24/593.htm'), (594, 846, 'जसबीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '1/3', 108, 15, '36K-5M-0S', 'dhani', 'https://taruana-archive.vercel.app/2023-24/594.htm'), (594, 847, 'जसबीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '1/3', None, 0, '0K-0M-0S', '', 'https://taruana-archive.vercel.app/2023-24/594.htm'), (595, 848, 'जसबीर सिंह', 'गुरबचन सिंह', 'बिशन सिंह', '1/3', 3, 12, '1K-4M-0S', 'pipline', 'https://taruana-archive.vercel.app/2023-24/595.htm')]

VILLAGE = "Taruaana"
HADBast = "315"
TEHSIL = "Kalanwali"
DISTRICT = "Sirsa"
JAMABANDI_YEAR = "2023–24"


def frac(value):
    return Fraction(str(value).strip())


def fmt_fraction(value):
    return str(value.numerator) if value.denominator == 1 else f"{value.numerator}/{value.denominator}"


def area(marlas):
    kanals, remaining = divmod(int(marlas), 20)
    return f"{kanals}K-{remaining}M"


class Main(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Jamabandi • Land Share Workspace")
        self.resize(1450, 900)
        self.url = ""
        self.rows = []
        self.build()
        self.load(int(self.k.currentText()))

    def build(self):
        root = QVBoxLayout()
        root.setContentsMargins(22, 20, 22, 18)
        root.setSpacing(12)

        central = QWidget()
        central.setLayout(root)
        self.setCentralWidget(central)

        hero = QFrame()
        hero.setObjectName("Hero")
        h = QHBoxLayout(hero)
        v = QVBoxLayout()
        title = QLabel("Jamabandi • Land Share Workspace")
        title.setObjectName("Title")
        subtitle = QLabel("Khewat-centric ownership, land-share analysis and source navigation")
        subtitle.setObjectName("Sub")
        v.addWidget(title)
        v.addWidget(subtitle)
        h.addLayout(v)
        h.addStretch()
        self.badge = QLabel("READY")
        self.badge.setObjectName("Badge")
        h.addWidget(self.badge)
        root.addWidget(hero)

        meta = QGroupBox("Village / Revenue Particulars")
        g = QGridLayout(meta)
        vals = [
            ("Village", VILLAGE), ("Hadbast No.", HADBast),
            ("Tehsil", TEHSIL), ("District", DISTRICT),
            ("Jamabandi Year", JAMABANDI_YEAR),
        ]
        for i, (label, value) in enumerate(vals):
            g.addWidget(QLabel(label), 0, i * 2)
            edit = QLineEdit(value)
            edit.setReadOnly(True)
            g.addWidget(edit, 0, i * 2 + 1)
        root.addWidget(meta)

        owner = QGroupBox("Owner Particulars")
        g = QGridLayout(owner)
        self.name = QLineEdit()
        self.father = QLineEdit()
        self.grand = QLineEdit()
        for label, edit, col in [
            ("Name", self.name, 0),
            ("Father's Name", self.father, 2),
            ("Grandfather's Name", self.grand, 4),
        ]:
            g.addWidget(QLabel(label), 0, col)
            g.addWidget(edit, 0, col + 1)
            edit.setReadOnly(True)
        root.addWidget(owner)

        controls = QHBoxLayout()
        controls.addWidget(QLabel("Khewat"))
        self.k = QComboBox()
        self.k.addItems(map(str, sorted(set(row[0] for row in DATA))))
        controls.addWidget(self.k)
        controls.addWidget(QLabel("Search"))
        self.search = QLineEdit()
        self.search.setPlaceholderText("Hissa, Khatoni, location or remarks")
        controls.addWidget(self.search, 1)
        self.export = QPushButton("Export Report")
        controls.addWidget(self.export)
        root.addLayout(controls)

        cards = QHBoxLayout()
        self.cv = {}
        for label in ["KHATONI(S)", "RECORDED LAND", "ENTRIES", "HISSA TOTAL"]:
            frame = QFrame()
            frame.setObjectName("Card")
            layout = QVBoxLayout(frame)
            small = QLabel(label)
            small.setObjectName("CL")
            value = QLabel("—")
            value.setObjectName("CV")
            layout.addWidget(small)
            layout.addWidget(value)
            self.cv[label] = value
            cards.addWidget(frame)
        root.addLayout(cards)

        splitter = QSplitter(Qt.Vertical)
        self.table = QTableWidget(0, 6)
        self.table.setHorizontalHeaderLabels(
            ["Khatoni", "Hissa", "Total Land", "Owner Share", "Detail / Location", "Remarks"]
        )
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.table.setAlternatingRowColors(True)
        self.table.verticalHeader().setVisible(False)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)

        header = self.table.horizontalHeader()
        for column in (0, 1, 2, 3):
            header.setSectionResizeMode(column, QHeaderView.ResizeToContents)
        header.setSectionResizeMode(4, QHeaderView.Stretch)
        header.setSectionResizeMode(5, QHeaderView.Stretch)
        splitter.addWidget(self.table)

        bottom = QWidget()
        bl = QHBoxLayout(bottom)

        source_box = QGroupBox("Source Nakal")
        source_layout = QVBoxLayout(source_box)
        self.src = QLabel("Select a row.")
        self.src.setWordWrap(True)
        self.src.setTextInteractionFlags(Qt.TextSelectableByMouse)
        self.open_button = QPushButton("Open Original Nakal")
        source_layout.addWidget(self.src)
        source_layout.addWidget(self.open_button)
        bl.addWidget(source_box, 2)

        analysis_box = QGroupBox("Share Analysis")
        analysis_layout = QVBoxLayout(analysis_box)
        self.analysis = QLabel("—")
        self.analysis.setWordWrap(True)
        analysis_layout.addWidget(self.analysis)
        bl.addWidget(analysis_box, 3)

        splitter.addWidget(bottom)
        splitter.setSizes([500, 190])
        root.addWidget(splitter, 1)

        self.k.currentTextChanged.connect(self._khewat_changed)
        self.search.textChanged.connect(self.filter)
        self.table.itemSelectionChanged.connect(self.select)
        self.open_button.clicked.connect(self.open_source)
        self.export.clicked.connect(self.export_xlsx)

        menu = self.menuBar().addMenu("&File")
        action = QAction("Export Excel / CSV", self)
        action.triggered.connect(self.export_xlsx)
        menu.addAction(action)

        action = QAction("Exit", self)
        action.setShortcut("Ctrl+Q")
        action.triggered.connect(self.close)
        menu.addAction(action)

        self.statusBar().showMessage("Ready")

    def _khewat_changed(self, value):
        if value:
            self.load(int(value))

    def load(self, khewat):
        self.rows = [row for row in DATA if row[0] == khewat]
        if not self.rows:
            return

        first = self.rows[0]
        self.name.setText(first[2])
        self.father.setText(first[3])
        self.grand.setText(first[4])

        self.cv["KHATONI(S)"].setText(", ".join(str(row[1]) for row in self.rows))
        self.cv["ENTRIES"].setText(str(len(self.rows)))

        recorded_marlas = sum((row[6] or 0) * 20 + (row[7] or 0) for row in self.rows)
        self.cv["RECORDED LAND"].setText(area(recorded_marlas))

        try:
            share_total = sum((frac(row[5]) for row in self.rows), Fraction())
            self.cv["HISSA TOTAL"].setText(fmt_fraction(share_total))
        except (ValueError, ZeroDivisionError):
            self.cv["HISSA TOTAL"].setText("—")

        self.table.setRowCount(0)
        for row in self.rows:
            table_row = self.table.rowCount()
            self.table.insertRow(table_row)
            values = [
                row[1], row[5],
                f"{row[6] or 0}K-{row[7] or 0}M",
                row[8], row[9], ""
            ]
            for column, value in enumerate(values):
                self.table.setItem(table_row, column, QTableWidgetItem(str(value)))
            self.table.item(table_row, 0).setData(Qt.UserRole, row[10])

        self.analysis.setText(
            f"<b>Khewat {khewat}</b><br>"
            f"Owner: {first[2]}<br>"
            f"Father: {first[3]}<br>"
            f"Grandfather: {first[4]}<br><br>"
            f"Recorded land: {area(recorded_marlas)} ({recorded_marlas} Marlas)<br>"
            f"Rows with zero/blank land are preserved as supplied."
        )

        self.url = ""
        self.src.setText("Select a row.")
        self.search.clear()
        self.statusBar().showMessage(f"Khewat {khewat} loaded • {len(self.rows)} entries")
        self.badge.setText("READY")

    def filter(self, text):
        needle = text.lower().strip()
        for row_index in range(self.table.rowCount()):
            values = [
                self.table.item(row_index, col).text().lower()
                for col in range(self.table.columnCount())
                if self.table.item(row_index, col) is not None
            ]
            haystack = " ".join(values)
            self.table.setRowHidden(row_index, bool(needle) and needle not in haystack)

    def select(self):
        selected = self.table.selectedItems()
        if not selected:
            return
        row_index = selected[0].row()
        row = self.rows[row_index]
        self.url = row[10]
        self.src.setText(f"Original Nakal:<br>{self.url}")
        self.analysis.setText(
            f"<b>Khewat {row[0]}</b><br>"
            f"Khatoni: {row[1]}<br>"
            f"Hissa: {row[5]}<br>"
            f"Owner Share: {row[8]}<br>"
            f"Detail / Location: {row[9] or '—'}"
        )

    def open_source(self):
        if not self.url:
            QMessageBox.information(self, "Source Nakal", "Please select a row first.")
            return
        webbrowser.open(self.url)

    def export_xlsx(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Export Report", "",
            "Excel Workbook (*.xlsx);;CSV File (*.csv)"
        )
        if not path or not self.rows:
            return
        try:
            if path.lower().endswith(".csv"):
                self._export_csv(path)
            else:
                self._export_excel(path)
            QMessageBox.information(self, "Export Complete", f"Saved:\n{path}")
            self.statusBar().showMessage(f"Exported report: {path}")
        except Exception as exc:
            QMessageBox.critical(self, "Export Error", str(exc))

    def _export_csv(self, path):
        with open(path, "w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(["Village", VILLAGE])
            writer.writerow(["Hadbast No.", HADBast])
            writer.writerow(["Tehsil", TEHSIL])
            writer.writerow(["District", DISTRICT])
            writer.writerow(["Jamabandi Year", JAMABANDI_YEAR])
            writer.writerow(["Khewat", self.rows[0][0]])
            writer.writerow(["Khatoni(s)", ", ".join(str(row[1]) for row in self.rows)])
            writer.writerow(["Owner", self.rows[0][2]])
            writer.writerow(["Father", self.rows[0][3]])
            writer.writerow(["Grandfather", self.rows[0][4]])
            writer.writerow([])
            writer.writerow(["Khatoni", "Hissa", "Total Land", "Owner Share", "Detail / Location", "Remarks"])
            for row in self.rows:
                writer.writerow([row[1], row[5], f"{row[6] or 0}K-{row[7] or 0}M", row[8], row[9], ""])

    def _export_excel(self, path):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment
            from openpyxl.worksheet.table import Table, TableStyleInfo
        except ImportError as exc:
            raise RuntimeError(
                "Excel export requires openpyxl. Install it with: pip install openpyxl"
            ) from exc

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Land Share Report"

        info = [
            ["Village", VILLAGE],
            ["Hadbast No.", HADBast],
            ["Tehsil", TEHSIL],
            ["District", DISTRICT],
            ["Jamabandi Year", JAMABANDI_YEAR],
            ["Khewat", self.rows[0][0]],
            ["Khatoni(s)", ", ".join(str(row[1]) for row in self.rows)],
            ["Name", self.rows[0][2]],
            ["Father's Name", self.rows[0][3]],
            ["Grandfather's Name", self.rows[0][4]],
            [],
        ]
        for line in info:
            sheet.append(line)

        sheet.append(["Khatoni", "Hissa", "Total Land", "Owner Share", "Detail / Location", "Remarks"])
        for row in self.rows:
            sheet.append([row[1], row[5], f"{row[6] or 0}K-{row[7] or 0}M", row[8], row[9], ""])

        for cell in sheet[12]:
            cell.font = Font(bold=True)

        sheet.freeze_panes = "A13"
        for column, width in zip("ABCDEF", [14, 20, 22, 25, 38, 25]):
            sheet.column_dimensions[column].width = width

        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        end_row = 12 + len(self.rows)
        table = Table(displayName="LandShares", ref=f"A12:F{end_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
        workbook.save(path)


STYLE = r"""
QMainWindow, QWidget { background: #f4f7fb; color: #172033; }
QMenuBar { background: white; padding: 5px; }
QFrame#Hero { background: #172033; border-radius: 16px; }
QLabel#Title { color: white; font-size: 26px; font-weight: 800; }
QLabel#Sub { color: #cbd5e1; font-size: 12px; }
QLabel#Badge { background: #e7f7ee; color: #176b3a; border-radius: 14px; padding: 8px 14px; font-weight: 800; }
QGroupBox { background: white; border: 1px solid #dbe2ec; border-radius: 12px; margin-top: 10px; font-weight: 700; padding-top: 12px; }
QGroupBox::title { subcontrol-origin: margin; left: 16px; padding: 0 6px; }
QLineEdit, QComboBox { background: #fbfcfe; border: 1px solid #cfd8e3; border-radius: 7px; padding: 8px; }
QPushButton { background: white; border: 1px solid #cfd8e3; border-radius: 8px; padding: 8px 14px; font-weight: 650; }
QPushButton:hover { background: #edf3fb; }
QTableWidget { background: white; border: 1px solid #dbe2ec; border-radius: 9px; gridline-color: #e8edf3; selection-background-color: #dce8f7; selection-color: #172033; }
QHeaderView::section { background: #eef2f7; border: 0; border-bottom: 1px solid #dbe2ec; padding: 9px; font-weight: 750; }
QFrame#Card { background: white; border: 1px solid #dbe2ec; border-radius: 10px; }
QLabel#CL { color: #697586; font-size: 10px; font-weight: 750; }
QLabel#CV { color: #172033; font-size: 18px; font-weight: 800; }
QStatusBar { background: white; }
"""


def main():
    app = QApplication(sys.argv)
    app.setApplicationName("Jamabandi Land Share Workspace")
    app.setStyleSheet(STYLE)
    window = Main()
    window.show()
    return app.exec()


if __name__ == "__main__":
    sys.exit(main())
